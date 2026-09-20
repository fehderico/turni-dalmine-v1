from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from typing import TYPE_CHECKING
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

if TYPE_CHECKING:
    from .scheduler import ScheduleResult


STATUS_COLORS = {
    "Completo": "C6E0B4",
    "Parziale": "FFE699",
    "Scoperto": "F4B084",
}

DAY_NAMES = ["Lunedì", "Martedì", "Mercoledì", "Giovedì", "Venerdì", "Sabato", "Domenica"]
PDF_FONT = "Helvetica"
PDF_FONT_BOLD = "Helvetica-Bold"


def _register_pdf_fonts() -> None:
    # I font standard PDF evitano differenze di incorporamento tra ambienti.
    return None


def build_frames(result: ScheduleResult) -> dict[str, pd.DataFrame]:
    assigned: dict[str, list[str]] = defaultdict(list)
    assigned_count: dict[str, int] = defaultdict(int)
    shift_by_id = {shift.id: shift for shift in result.shifts}
    for doctor_key, shift_id in result.assignments:
        doctor = result.doctors[doctor_key]
        shift = shift_by_id[shift_id]
        assigned[shift_id].append(doctor.display_name)
        assigned_count[doctor_key] += 1

    calendar_rows = []
    shortage_rows = []
    max_slots = max((s.definition.required for s in result.shifts), default=0)
    for shift in result.shifts:
        names = sorted(assigned[shift.id])
        candidates = result.candidate_counts.get(shift.id, 0)
        uncovered = shift.definition.required - len(names)
        status = "Completo" if uncovered == 0 else ("Scoperto" if not names else "Parziale")
        row = {
            "Data": shift.date,
            "Giorno": DAY_NAMES[shift.date.weekday()],
            "Codice": shift.definition.code,
            "Turno": shift.definition.label,
            "Orario": f"{shift.start:%H:%M}-{shift.end:%H:%M}",
            "Richiesti": shift.definition.required,
            "Disponibili": candidates,
            "Assegnati": len(names),
            "Scoperti": uncovered,
            "Stato": status,
        }
        row.update({f"Medico {i + 1}": names[i] if i < len(names) else "" for i in range(max_slots)})
        calendar_rows.append(row)
        if uncovered:
            if candidates == 0:
                reason = "Nessuna disponibilità dichiarata"
            elif candidates < shift.definition.required:
                reason = (
                    "Solo un medico ha dichiarato disponibilità"
                    if candidates == 1
                    else f"Solo {candidates} medici hanno dichiarato disponibilità"
                )
            else:
                reason = f"{candidates} disponibili; riposi o sovrapposizioni riducono le assegnazioni compatibili"
            shortage_rows.append({
                "Data": shift.date,
                "Codice": shift.definition.code,
                "Turno": shift.definition.label,
                "Richiesti": shift.definition.required,
                "Coperti": len(names),
                "Scoperti": uncovered,
                "Disponibili dichiarati": candidates,
                "Motivo": reason,
            })

    assignment_rows = []
    assignments_by_doctor: dict[str, list[str]] = defaultdict(list)
    for doctor_key, shift_id in result.assignments:
        shift = shift_by_id[shift_id]
        doctor = result.doctors[doctor_key]
        assignments_by_doctor[doctor_key].append(f"{shift.date:%d/%m} {shift.definition.code}")
        assignment_rows.append({
            "Data": shift.date,
            "Giorno": DAY_NAMES[shift.date.weekday()],
            "Codice": shift.definition.code,
            "Turno": shift.definition.label,
            "Orario": f"{shift.start:%H:%M}-{shift.end:%H:%M}",
            "Medico": doctor.display_name,
            "ASST": doctor.ats,
            "Email": doctor.email,
        })

    doctor_rows = []
    for key, doctor in sorted(result.doctors.items(), key=lambda item: item[1].display_name):
        doctor_rows.append({
            "Nome": doctor.nome,
            "Cognome": doctor.cognome,
            "Email": doctor.email,
            "ASST": doctor.ats,
            "Stato contratto": doctor.contract_status,
            "Turni assegnati": assigned_count[key],
            "Dettaglio turni": ", ".join(sorted(assignments_by_doctor[key])),
            "Note": doctor.notes,
        })

    parameters = [
        ("Mese", result.month),
        ("Riposo minimo", "11 ore"),
        ("Priorità", "Copertura > Bergamo Ovest > Bergamo Est > altri > estrazione casuale riproducibile"),
        ("Minimi contrattuali", "Non utilizzati nella V1"),
        ("Codice estrazione", result.seed),
        ("Formato origine", result.source_format),
        ("Foglio origine", result.source_sheet),
        ("Identità", "Email per Google Forms; nome + cognome nel formato normalizzato"),
        ("Confine mese", "Non gestito in v1"),
    ]
    for phase, status in result.phase_status.items():
        parameters.append((f"Esito {phase}", status))

    return {
        "Calendario": pd.DataFrame(calendar_rows),
        "Assegnazioni": pd.DataFrame(assignment_rows),
        "Riepilogo medici": pd.DataFrame(doctor_rows),
        "Scoperture": pd.DataFrame(shortage_rows),
        "Controlli": pd.DataFrame({"Avviso": result.warnings or ["Nessun avviso"]}),
        "Parametri": pd.DataFrame(parameters, columns=["Parametro", "Valore"]),
}


def _style_excel_sheet(sheet, frame: pd.DataFrame) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, column in enumerate(frame.columns, start=1):
        values = [str(column)] + [str(v) for v in frame[column].fillna("").tolist()]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(v) for v in values) + 2, 45)
    for column_index, column in enumerate(frame.columns, start=1):
        if column == "Data":
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column_index).number_format = "dd/mm/yyyy"
    if "Stato" in frame.columns and not frame.empty:
        status_col = list(frame.columns).index("Stato") + 1
        for row in range(2, sheet.max_row + 1):
            status = sheet.cell(row, status_col).value
            fill = STATUS_COLORS.get(status)
            if fill:
                sheet.cell(row, status_col).fill = PatternFill("solid", fgColor=fill)
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def export_excel(result: ScheduleResult) -> bytes:
    frames = build_frames(result)
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, frame in frames.items():
            frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)
            sheet = writer.book[sheet_name[:31]]
            _style_excel_sheet(sheet, frame)
    return output.getvalue()


def export_frame_excel(frame: pd.DataFrame, sheet_name: str = "Vista") -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)
        _style_excel_sheet(writer.book[sheet_name[:31]], frame)
    return output.getvalue()


def _display_value(value: object) -> str:
    if pd.isna(value):
        return ""
    if hasattr(value, "strftime"):
        try:
            return value.strftime("%d/%m/%Y")
        except (TypeError, ValueError):
            pass
    return str(value)


def export_frame_pdf(frame: pd.DataFrame, title: str) -> bytes:
    _register_pdf_fonts()
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=title,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "FrameTitle", parent=styles["Title"], alignment=TA_CENTER,
        textColor=colors.HexColor("#17365D"), fontName=PDF_FONT_BOLD,
    )
    cell_style = ParagraphStyle("FrameCell", parent=styles["BodyText"], fontName=PDF_FONT, fontSize=6.5, leading=7.5)
    header_style = ParagraphStyle("FrameHeader", parent=cell_style, fontName=PDF_FONT_BOLD, textColor=colors.white)
    safe_frame = frame.copy()
    if safe_frame.empty:
        safe_frame = pd.DataFrame({"Risultato": ["Nessun dato per i filtri selezionati"]})
    rows = [[Paragraph(escape(str(column)), header_style) for column in safe_frame.columns]]
    for _, row in safe_frame.iterrows():
        rows.append([Paragraph(escape(_display_value(value)), cell_style) for value in row])
    available_width = landscape(A4)[0] - 20 * mm
    weights = []
    for column in safe_frame.columns:
        sample = [len(str(column))] + [min(len(_display_value(v)), 45) for v in safe_frame[column].head(80)]
        weights.append(max(7, max(sample)))
    total_weight = sum(weights) or 1
    widths = [available_width * weight / total_weight for weight in weights]
    table = Table(rows, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF4FA")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    doc.build([Paragraph(escape(title), title_style), Spacer(1, 5 * mm), table])
    return output.getvalue()


def export_pdf(result: ScheduleResult) -> bytes:
    _register_pdf_fonts()
    frames = build_frames(result)
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=f"Calendario turni {result.month}",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("CenterTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#17365D"), fontName=PDF_FONT_BOLD)
    heading_style = ParagraphStyle("Heading", parent=styles["Heading1"], textColor=colors.HexColor("#17365D"), fontName=PDF_FONT_BOLD)
    tiny = ParagraphStyle("Tiny", parent=styles["BodyText"], fontName=PDF_FONT, fontSize=6.7, leading=8)
    story = [Paragraph(f"Calendario turni centrale di Dalmine - {escape(result.month)}", title_style), Spacer(1, 5 * mm)]

    calendar = frames["Calendario"]
    for start in range(0, len(calendar), 18):
        part = calendar.iloc[start:start + 18]
        rows = [["Data", "Turno", "Orario", "Copertura", "Medici"]]
        for _, row in part.iterrows():
            doctors = [str(row[c]) for c in calendar.columns if c.startswith("Medico ") and str(row[c]).strip()]
            rows.append([
                row["Data"].strftime("%d/%m/%Y"),
                f"{row['Codice']} — {row['Turno']}",
                row["Orario"],
                f"{row['Assegnati']}/{row['Richiesti']} ({row['Stato']})",
                Paragraph(escape(", ".join(doctors) or "—"), tiny),
            ])
        table = Table(rows, colWidths=[25 * mm, 35 * mm, 28 * mm, 37 * mm, 142 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), PDF_FONT_BOLD),
            ("FONTNAME", (0, 1), (-1, -1), PDF_FONT),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF4FA")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(table)
        if start + 18 < len(calendar):
            story.append(PageBreak())

    story.extend([PageBreak(), Paragraph("Scoperture", heading_style)])
    shortage = frames["Scoperture"]
    shortage_rows = [["Data", "Turno", "Coperti", "Scoperti"]]
    for _, row in shortage.iterrows():
        shortage_rows.append([row["Data"].strftime("%d/%m/%Y"), row["Turno"], row["Coperti"], row["Scoperti"]])
    if len(shortage_rows) == 1:
        shortage_rows.append(["—", "Nessuna scopertura", "—", "—"])
    shortage_table = Table(shortage_rows, colWidths=[35 * mm, 70 * mm, 30 * mm, 30 * mm], repeatRows=1)
    shortage_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), PDF_FONT_BOLD),
        ("FONTNAME", (0, 1), (-1, -1), PDF_FONT),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")),
        ("ALIGN", (2, 1), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    story.append(shortage_table)

    story.extend([PageBreak(), Paragraph("Turni per medico", heading_style)])
    doctors = frames["Riepilogo medici"]
    doctor_rows = [["Medico", "ASST", "Turni", "Dettaglio"]]
    for _, row in doctors.iterrows():
        doctor_rows.append([
            f"{row['Nome']} {row['Cognome']}",
            row["ASST"],
            row["Turni assegnati"],
            Paragraph(escape(str(row["Dettaglio turni"]) or "—"), tiny),
        ])
    doctor_table = Table(doctor_rows, colWidths=[55 * mm, 45 * mm, 25 * mm, 142 * mm], repeatRows=1)
    doctor_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), PDF_FONT_BOLD),
        ("FONTNAME", (0, 1), (-1, -1), PDF_FONT),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF4FA")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
    ]))
    story.append(doctor_table)
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(PDF_FONT, 7)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(10 * mm, 6 * mm, f"Centrale di Dalmine - calendario {result.month}")
        canvas.drawRightString(landscape(A4)[0] - 10 * mm, 6 * mm, f"Pagina {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
